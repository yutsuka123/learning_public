"""
python-docx / python-pptxによるWord・PowerPoint自動化サンプル。

概要:
    `.docx`（Word）・`.pptx`（PowerPoint）ファイルを、それぞれのアプリを起動せずに
    Pythonだけで生成する。定型帳票・報告書テンプレートの自動生成等に使う。
主な仕様:
    - demonstrateWordDocument(): 見出し・段落・表を含むWord文書を作成する。
    - demonstratePowerPointSlides(): タイトルスライド・箇条書きスライドを作成する。
制限事項:
    - どちらもExcel同様、`.docx`/`.pptx`（Office 2007以降のXMLベース形式）専用。

実行方法:
    cd python/office_automation
    pip install python-docx python-pptx
    python3 word_powerpoint_automation.py
    # 検証環境: python-docx, python-pptx で実行結果を確認済み。
"""

import os

from docx import Document
from pptx import Presentation


def demonstrateWordDocument():
    """
    見出し・段落・表を含むWord文書を作成する。

    実行結果（例）:
        saved report.docx (36833 bytes)
    """

    print("=== 1. Word document: heading, paragraph, table ===")

    outputDir = os.path.dirname(__file__) or "."
    document = Document()

    document.add_heading("Monthly Report", level=1)
    document.add_paragraph("This report summarizes sensor readings for the month.")

    document.add_heading("Readings", level=2)
    table = document.add_table(rows=1, cols=2)
    headerCells = table.rows[0].cells
    headerCells[0].text = "Sensor"
    headerCells[1].text = "Value"
    for sensorName, value in [("temperature", "25.5"), ("humidity", "60.0")]:
        rowCells = table.add_row().cells
        rowCells[0].text = sensorName
        rowCells[1].text = value

    path = os.path.join(outputDir, "report.docx")
    document.save(path)
    print(f"saved report.docx ({os.path.getsize(path)} bytes)")


def demonstratePowerPointSlides():
    """
    タイトルスライドと箇条書きスライドを作成する。

    実行結果（例）:
        saved slides.pptx (29170 bytes), slide count=2
    """

    print("\n=== 2. PowerPoint slides: title + bullet list ===")

    outputDir = os.path.dirname(__file__) or "."
    presentation = Presentation()

    # スライド1: タイトルスライド（既定レイアウト0番）
    titleSlideLayout = presentation.slide_layouts[0]
    slide1 = presentation.slides.add_slide(titleSlideLayout)
    slide1.shapes.title.text = "Monthly Report"
    slide1.placeholders[1].text = "Sensor data summary"

    # スライド2: 箇条書き（既定レイアウト1番=タイトル+コンテンツ）
    bulletSlideLayout = presentation.slide_layouts[1]
    slide2 = presentation.slides.add_slide(bulletSlideLayout)
    slide2.shapes.title.text = "Key Findings"
    body = slide2.placeholders[1].text_frame
    body.text = "Temperature stable around 25.5C"
    for bulletText in ["Humidity averaged 60%", "No sensor faults detected"]:
        paragraph = body.add_paragraph()
        paragraph.text = bulletText
        paragraph.level = 1

    path = os.path.join(outputDir, "slides.pptx")
    presentation.save(path)
    slideCount = len(list(presentation.slides))
    print(f"saved slides.pptx ({os.path.getsize(path)} bytes), slide count={slideCount}")


if __name__ == "__main__":
    demonstrateWordDocument()
    demonstratePowerPointSlides()
