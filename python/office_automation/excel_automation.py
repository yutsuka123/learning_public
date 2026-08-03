"""
openpyxlによるExcel自動化サンプル。

概要:
    Excelファイル(.xlsx)の作成・書き込み・数式・書式設定・読み込みを、Excel本体を
    起動せずにPythonだけで行う。定型レポートの自動生成や、大量データの一括転記等に使う。
主な仕様:
    - demonstrateWriteBasics(): セルへの書き込み、行の一括追加。
    - demonstrateFormulasAndFormatting(): Excel数式の埋め込み、セルの書式設定。
    - demonstrateReadBack(): 書き込んだファイルを読み込み、値を確認する。
    - demonstrateFromPandas(): pandasのDataFrameをExcelシートへ変換する。
制限事項:
    - openpyxlは`.xlsx`（Excel 2007以降の形式）専用。古い`.xls`形式は非対応。
    - 数式は「文字列として書き込む」だけで、Python側では計算しない
      （実際に開いたExcelやLibreOffice等が計算する）。

実行方法:
    cd python/office_automation
    pip install openpyxl pandas
    python3 excel_automation.py
    # 検証環境: openpyxl 3.1.5 で実行結果を確認済み。
"""

import os

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def demonstrateWriteBasics():
    """
    セルへの書き込みと、行の一括追加を確認する。

    実行結果（例）:
        A1=title
        B1=42
        after append: 4 rows (including header)
        saved basic_report.xlsx (4902 bytes)
    """

    print("=== 1. write basics: cells, append ===")

    outputDir = os.path.dirname(__file__) or "."
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"

    sheet["A1"] = "title"
    sheet["B1"] = 42
    print(f"A1={sheet['A1'].value}")
    print(f"B1={sheet['B1'].value}")

    sheet.append(["name", "score"])  # ヘッダー行
    sheet.append(["alice", 90])
    sheet.append(["bob", 75])
    print(f"after append: {sheet.max_row} rows (including header)")

    path = os.path.join(outputDir, "basic_report.xlsx")
    workbook.save(path)
    print(f"saved basic_report.xlsx ({os.path.getsize(path)} bytes)")


def demonstrateFormulasAndFormatting():
    """
    Excel数式の埋め込みと、セルの書式設定（太字・背景色・列幅）を確認する。

    実行結果（例）:
        saved formulas_report.xlsx (5009 bytes)
        C2 formula string (Excel側で計算される): =A2*B2
    """

    print("\n=== 2. formulas & formatting ===")

    outputDir = os.path.dirname(__file__) or "."
    workbook = Workbook()
    sheet = workbook.active

    headers = ["quantity", "unit_price", "subtotal"]
    sheet.append(headers)
    for cell in sheet[1]:
        # [重要] Font/PatternFillでセルの見た目を制御する。Excelを手作業で操作するのと
        # 同じことをコードで再現できる。
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

    sheet.append([3, 100])
    sheet.append([5, 200])

    # [重要] 数式は"="で始まる文字列としてそのまま書き込む。Python側は計算しない
    # （セルの値としてExcel/LibreOffice等が開いたときに初めて計算される）。
    for rowIndex in range(2, sheet.max_row + 1):
        sheet[f"C{rowIndex}"] = f"=A{rowIndex}*B{rowIndex}"

    for columnIndex in range(1, 4):
        sheet.column_dimensions[get_column_letter(columnIndex)].width = 14

    path = os.path.join(outputDir, "formulas_report.xlsx")
    workbook.save(path)
    print(f"saved formulas_report.xlsx ({os.path.getsize(path)} bytes)")
    print(f"C2 formula string (Excel側で計算される): {sheet['C2'].value}")


def demonstrateReadBack():
    """
    書き込んだファイルを読み込み、値を確認する。

    実行結果（例）:
        read back: name=alice, score=90
        read back: name=bob, score=75
    """

    print("\n=== 3. read back an existing .xlsx ===")

    outputDir = os.path.dirname(__file__) or "."
    path = os.path.join(outputDir, "basic_report.xlsx")

    workbook = load_workbook(path)
    sheet = workbook["Report"]

    # [重要] basic_report.xlsxの行構成: 1行目=("title","42")、2行目=("name","score")ヘッダー、
    # 3〜4行目=データ。実際に作成時、min_row=2を指定してしまい"name/score"ヘッダー行まで
    # 読み込んでしまう間違いをこの場で発見・修正した（min_row=3が正しい）。
    for row in sheet.iter_rows(min_row=3, values_only=True):
        name, score = row
        print(f"read back: name={name}, score={score}")


def demonstrateFromPandas():
    """
    pandasのDataFrameをExcelシートへ変換する。

    目的:
        `python/data_analysis/pandas_basics.py`で扱ったDataFrameを、そのまま
        Excelレポートとして出力する、実務で頻出のパターンを確認する。

    実行結果（例）:
        saved from_pandas.xlsx (4897 bytes)
        sheet names: ['Sales']
    """

    print("\n=== 4. DataFrame -> Excel (pandas + openpyxl) ===")

    outputDir = os.path.dirname(__file__) or "."
    path = os.path.join(outputDir, "from_pandas.xlsx")

    df = pd.DataFrame(
        {
            "region": ["east", "west", "east"],
            "amount": [100, 150, 200],
        }
    )

    # [重要] pandasのExcel書き込みは内部でopenpyxl（または他エンジン）を使う。
    # engine="openpyxl"を明示すると、このファイルの他の例と同じエンジンで統一できる。
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sales", index=False)

    print(f"saved from_pandas.xlsx ({os.path.getsize(path)} bytes)")

    workbook = load_workbook(path)
    print(f"sheet names: {workbook.sheetnames}")


if __name__ == "__main__":
    demonstrateWriteBasics()
    demonstrateFormulasAndFormatting()
    demonstrateReadBack()
    demonstrateFromPandas()
